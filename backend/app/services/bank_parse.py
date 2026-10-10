"""Bank statement file parsing (CSV, XLSX, legacy XLS) into normalized rows.

The importer never assumes a fixed layout: it finds the header row, guesses the
column mapping from known header names (HDFC and common variants), and lets the
user override everything. All amounts are parsed exactly into paise.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import re
from dataclasses import dataclass, field
from typing import Any

from ..errors import DomainError
from ..money import MoneyParseError, parse_rupees

FIELDS = ["date", "narration", "reference", "value_date", "debit", "credit", "balance", "amount", "dr_cr"]
REQUIRED_ONE_OF = [("date",), ("narration",), ("debit", "credit", "amount")]

SYNONYMS: dict[str, list[str]] = {
    "date": ["date", "txn date", "transaction date", "tran date", "posting date", "trans date"],
    "narration": ["narration", "description", "particulars", "remarks", "transaction details", "details",
                  "transaction remarks"],
    "reference": ["chq./ref.no.", "chq/ref no", "chq / ref no", "ref no", "ref no./cheque no.", "reference",
                  "reference no", "cheque no", "chq no", "chq./ref. no.", "utr", "utr no", "ref. no."],
    "value_date": ["value dt", "value date"],
    "debit": ["withdrawal amt.", "withdrawal amt", "withdrawal", "withdrawals", "debit", "debit amount",
              "debit amt", "dr amount", "withdrawal amount"],
    "credit": ["deposit amt.", "deposit amt", "deposit", "deposits", "credit", "credit amount", "credit amt",
               "cr amount", "deposit amount"],
    "balance": ["closing balance", "balance", "running balance", "closing bal", "balance amt"],
    "amount": ["amount", "transaction amount", "txn amount"],
    "dr_cr": ["dr/cr", "cr/dr", "type", "dr / cr", "debit/credit"],
}

MAX_ROWS = 20000


def _norm_header(v: Any) -> str:
    return re.sub(r"\s+", " ", str(v or "").strip().lower())


@dataclass
class Table:
    sheets: list[str]
    sheet: str | None
    rows: list[list[Any]]


def read_table(data: bytes, content_type: str, sheet: str | None = None) -> Table:
    if content_type == "text/csv":
        return Table([], None, _read_csv(data))
    if content_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
        return _read_xlsx(data, sheet)
    if content_type == "application/vnd.ms-excel":
        return _read_xls(data, sheet)
    raise DomainError("unsupported_file", "Statement must be a CSV, XLSX or XLS file")


def _read_csv(data: bytes) -> list[list[Any]]:
    for enc in ("utf-8-sig", "cp1252"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:  # pragma: no cover - sniff_type already rejects these
        raise DomainError("bad_encoding", "Could not read the CSV text encoding")
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = [list(r) for r in csv.reader(io.StringIO(text), dialect)]
    return rows[:MAX_ROWS]


def _read_xlsx(data: bytes, sheet: str | None) -> Table:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises many types for corrupt files
        raise DomainError("bad_file", "The XLSX file could not be opened. Is it a valid Excel file?") from exc
    try:
        names = wb.sheetnames
        name = sheet if sheet in names else names[0]
        ws = wb[name]
        rows = []
        for i, r in enumerate(ws.iter_rows(values_only=True)):
            if i >= MAX_ROWS:
                break
            rows.append(list(r))
        return Table(names, name, rows)
    finally:
        wb.close()


def _read_xls(data: bytes, sheet: str | None) -> Table:
    import xlrd

    try:
        book = xlrd.open_workbook(file_contents=data)
    except Exception as exc:
        raise DomainError("bad_file", "The XLS file could not be opened. Try exporting the statement as XLSX or CSV.") from exc
    names = book.sheet_names()
    name = sheet if sheet in names else names[0]
    sh = book.sheet_by_name(name)
    rows = []
    for ri in range(min(sh.nrows, MAX_ROWS)):
        row = []
        for ci in range(sh.ncols):
            cell = sh.cell(ri, ci)
            if cell.ctype == xlrd.XL_CELL_DATE:
                row.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode))
            elif cell.ctype == xlrd.XL_CELL_NUMBER:
                v = cell.value
                row.append(int(v) if float(v).is_integer() else v)
            else:
                row.append(cell.value)
        rows.append(row)
    return Table(names, name, rows)


def _match_field(header: str) -> str | None:
    h = _norm_header(header)
    if not h:
        return None
    for fld, names in SYNONYMS.items():
        if h in names:
            return fld
    # Loose match (e.g. "Withdrawal Amt.(INR)")
    for fld, names in SYNONYMS.items():
        for n in names:
            if len(n) >= 5 and h.startswith(n):
                return fld
    return None


def detect_header(rows: list[list[Any]]) -> int | None:
    best, best_score = None, 0
    for i, row in enumerate(rows[:80]):
        fields = {_match_field(c) for c in row} - {None}
        score = len(fields)
        if "date" in fields and ({"debit", "credit", "amount"} & fields) and score > best_score:
            best, best_score = i, score
    return best if best_score >= 3 else None


def auto_mapping(rows: list[list[Any]]) -> dict:
    header_row = detect_header(rows)
    columns: dict[str, int] = {}
    if header_row is not None:
        for idx, cell in enumerate(rows[header_row]):
            fld = _match_field(cell)
            if fld and fld not in columns:
                columns[fld] = idx
    return {"header_row": header_row, "columns": columns, "date_format": "auto"}


def validate_mapping(mapping: dict, ncols: int) -> None:
    cols = mapping.get("columns") or {}
    for k, v in cols.items():
        if k not in FIELDS:
            raise DomainError("invalid_mapping", f"Unknown field '{k}'")
        if not isinstance(v, int) or not 0 <= v < max(ncols, 1):
            raise DomainError("invalid_mapping", f"Column for '{k}' is out of range")
    if "date" not in cols or "narration" not in cols:
        raise DomainError("invalid_mapping", "Choose the columns for Date and Narration")
    if not ({"debit", "credit"} <= cols.keys() or "amount" in cols):
        raise DomainError("invalid_mapping", "Choose both Withdrawal and Deposit columns, or a single Amount column")
    hr = mapping.get("header_row")
    if hr is not None and (not isinstance(hr, int) or hr < 0):
        raise DomainError("invalid_mapping", "Invalid header row")


DATE_FORMATS = ["%d/%m/%y", "%d/%m/%Y", "%d-%m-%Y", "%d-%m-%y", "%Y-%m-%d", "%d %b %Y", "%d-%b-%Y", "%d-%b-%y",
                "%d %b %y", "%d.%m.%Y", "%d.%m.%y", "%d %B %Y", "%Y-%m-%d %H:%M:%S", "%d/%m/%Y %H:%M:%S"]


def parse_date(value: Any, fmt: str = "auto") -> dt.date | None:
    if value is None or value == "":
        return None
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    text = str(value).strip()
    if not text:
        return None
    formats = DATE_FORMATS if fmt in (None, "", "auto") else [fmt]
    for f in formats:
        try:
            return dt.datetime.strptime(text, f).date()
        except ValueError:
            continue
    raise ValueError(f"Unrecognised date '{text[:30]}'")


@dataclass
class ParsedRow:
    row_number: int  # 1-based, as seen in a spreadsheet
    txn_date: dt.date | None = None
    value_date: dt.date | None = None
    narration: str = ""
    reference: str | None = None
    debit_paise: int = 0
    credit_paise: int = 0
    balance_paise: int | None = None
    errors: list[str] = field(default_factory=list)


def _cell(row: list[Any], idx: int | None) -> Any:
    if idx is None or idx >= len(row):
        return None
    v = row[idx]
    return v.strip() if isinstance(v, str) else v


def _is_blank(v: Any) -> bool:
    return v is None or (isinstance(v, str) and v.strip() in {"", "-"})


def _clean_reference(v: Any) -> str | None:
    if _is_blank(v):
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    if not s or set(s) <= {"0"}:
        return None
    return s[:120]


def parse_rows(rows: list[list[Any]], mapping: dict) -> list[ParsedRow]:
    cols = mapping["columns"]
    fmt = mapping.get("date_format") or "auto"
    header_row = mapping.get("header_row")
    start = 0 if header_row is None else header_row + 1  # header on row 0 is common in CSV exports
    out: list[ParsedRow] = []
    for i in range(start, len(rows)):
        row = rows[i]
        raw_date = _cell(row, cols.get("date"))
        amounts_raw = [_cell(row, cols.get(k)) for k in ("debit", "credit", "amount")]
        narr = _cell(row, cols.get("narration"))
        # Skip separators / blank lines / footers that have no date and no amounts.
        if _is_blank(raw_date) and all(_is_blank(a) for a in amounts_raw):
            continue
        if isinstance(raw_date, str) and set(raw_date) <= {"*", "-", "=", " "}:
            continue
        pr = ParsedRow(row_number=i + 1)
        try:
            pr.txn_date = parse_date(raw_date, fmt)
        except ValueError as e:
            # After real rows, a text "date" without digits (e.g. "STATEMENT SUMMARY") marks the footer.
            if out and isinstance(raw_date, str) and not re.search(r"\d", raw_date):
                break
            pr.errors.append(str(e))
        if pr.txn_date is None and not pr.errors:
            pr.errors.append("Missing date")
        try:
            vd = _cell(row, cols.get("value_date"))
            pr.value_date = parse_date(vd, fmt) if not _is_blank(vd) else None
        except ValueError:
            pr.value_date = None  # value date is informational only
        pr.narration = re.sub(r"\s+", " ", str(narr or "")).strip()[:2000]
        pr.reference = _clean_reference(_cell(row, cols.get("reference")))
        try:
            if "amount" in cols:
                amt = parse_rupees(_cell(row, cols["amount"])) or 0
                ind = str(_cell(row, cols.get("dr_cr")) or "").strip().upper()
                if ind.startswith("D") or (not ind and amt < 0):
                    pr.debit_paise = abs(amt)
                elif ind.startswith("C") or (not ind and amt > 0):
                    pr.credit_paise = abs(amt)
            else:
                pr.debit_paise = parse_rupees(_cell(row, cols.get("debit"))) or 0
                pr.credit_paise = parse_rupees(_cell(row, cols.get("credit"))) or 0
            if "balance" in cols:
                pr.balance_paise = parse_rupees(_cell(row, cols["balance"]))
        except MoneyParseError as e:
            pr.errors.append(str(e))
        if pr.debit_paise < 0 or pr.credit_paise < 0:
            pr.errors.append("Negative amount in a withdrawal/deposit column")
        if pr.debit_paise and pr.credit_paise:
            pr.errors.append("Row has both a withdrawal and a deposit")
        if not pr.errors and not pr.debit_paise and not pr.credit_paise:
            pr.errors.append("Row has no amount")
        if not pr.narration and not pr.errors:
            pr.errors.append("Missing narration")
        out.append(pr)
    return out


# ------------------------------------------------------------------ statement self-checks
#
# Banks print their own summary (period, opening/closing balance, debit/credit totals and
# counts) around the transaction table. Reading it lets the importer prove that every row was
# captured. Labels are matched generically; nothing about a particular statement is hardcoded.

_PERIOD_RE = re.compile(
    r"from\s*:?\s*(\d{1,2}[/-]\d{1,2}[/-]\d{2,4})\s*to\s*:?\s*(\d{1,2}[/-]\d{1,2}[/-]\d{2,4})", re.I)
SUMMARY_LABELS = {
    "opening balance": "opening_paise", "opening bal": "opening_paise",
    "closing balance": "closing_paise", "closing bal": "closing_paise",
    "debits": "debit_total_paise", "total debits": "debit_total_paise", "withdrawals": "debit_total_paise",
    "credits": "credit_total_paise", "total credits": "credit_total_paise", "deposits": "credit_total_paise",
    "dr count": "debit_count", "debit count": "debit_count", "cr count": "credit_count", "credit count": "credit_count",
}


@dataclass
class StatementInfo:
    period_start: dt.date | None = None
    period_end: dt.date | None = None
    opening_paise: int | None = None
    closing_paise: int | None = None
    debit_total_paise: int | None = None
    credit_total_paise: int | None = None
    debit_count: int | None = None
    credit_count: int | None = None

    def as_dict(self) -> dict:
        return {k: (v.isoformat() if isinstance(v, dt.date) else v) for k, v in self.__dict__.items()}


def _parse_any_date(text: str) -> dt.date | None:
    for f in ("%d/%m/%Y", "%d/%m/%y", "%d-%m-%Y", "%d-%m-%y"):
        try:
            return dt.datetime.strptime(text, f).date()
        except ValueError:
            continue
    return None


def extract_statement_info(rows: list[list[Any]], header_row: int | None, table_end: int) -> StatementInfo:
    """Read the stated period (before the table) and the summary block (after the table).

    `table_end` is the 0-based index just past the last transaction row, so the
    transaction header's own "Closing Balance" column is never mistaken for the summary.
    """
    info = StatementInfo()
    for row in rows[: header_row if header_row is not None else 80]:
        for c in row:
            if isinstance(c, str):
                m = _PERIOD_RE.search(c)
                if m:
                    info.period_start, info.period_end = _parse_any_date(m.group(1)), _parse_any_date(m.group(2))
                    break
        if info.period_start:
            break
    for i in range(table_end, len(rows) - 1):
        for j, c in enumerate(rows[i]):
            key = SUMMARY_LABELS.get(_norm_header(c).rstrip(":").strip())
            if not key or getattr(info, key) is not None:
                continue
            below = rows[i + 1][j] if j < len(rows[i + 1]) else None
            try:
                if key.endswith("_count"):
                    if isinstance(below, (int, float)) and float(below).is_integer():
                        setattr(info, key, int(below))
                    elif isinstance(below, str) and below.strip().isdigit():
                        setattr(info, key, int(below.strip()))
                else:
                    setattr(info, key, parse_rupees(below))
            except MoneyParseError:
                pass
    return info


def _check(code: str, status: str, message: str, **extra) -> dict:
    return {"code": code, "status": status, "message": message, **extra}


def statement_checks(parsed: list[ParsedRow], info: StatementInfo) -> list[dict]:
    """Cross-check parsed rows against the bank's own summary. status: ok | warn | fail | info."""
    good = [r for r in parsed if not r.errors]
    checks: list[dict] = []
    bad = len(parsed) - len(good)
    if bad:
        checks.append(_check("unreadable_rows", "fail", f"{bad} rows could not be read."))
    debits = [r for r in good if r.debit_paise]
    credits = [r for r in good if r.credit_paise]
    if info.debit_count is not None and info.credit_count is not None:
        ok = (len(debits), len(credits)) == (info.debit_count, info.credit_count)
        checks.append(_check(
            "counts", "ok" if ok else "fail",
            f"Statement summary lists {info.debit_count} debits and {info.credit_count} credits; "
            f"read {len(debits)} debits and {len(credits)} credits."))
    if info.debit_total_paise is not None and info.credit_total_paise is not None:
        d, c = sum(r.debit_paise for r in debits), sum(r.credit_paise for r in credits)
        ok = (d, c) == (info.debit_total_paise, info.credit_total_paise)
        checks.append(_check("totals", "ok" if ok else "fail",
                             "Debit and credit totals match the statement summary." if ok else
                             "Debit/credit totals differ from the statement summary."))
    if None not in (info.opening_paise, info.closing_paise, info.debit_total_paise, info.credit_total_paise):
        ok = info.opening_paise - info.debit_total_paise + info.credit_total_paise == info.closing_paise
        checks.append(_check("balance_equation", "ok" if ok else "warn",
                             "Opening balance - debits + credits = closing balance." if ok else
                             "The statement summary's own balances do not add up."))
    if good and all(r.balance_paise is not None for r in good):
        breaks = [b.row_number for a, b in zip(good, good[1:])
                  if a.balance_paise - b.debit_paise + b.credit_paise != b.balance_paise]
        first = good[0]
        if info.opening_paise is not None and \
                info.opening_paise - first.debit_paise + first.credit_paise != first.balance_paise:
            breaks.insert(0, first.row_number)
        if info.closing_paise is not None and good[-1].balance_paise != info.closing_paise:
            breaks.append(good[-1].row_number)
        checks.append(_check(
            "running_balance", "ok" if not breaks else "fail",
            "Every row's closing balance follows from the previous one." if not breaks else
            f"Closing balance does not follow on at rows {', '.join(map(str, breaks[:8]))}; rows may be missing "
            "or mis-read.", rows=breaks[:50]))
    if info.period_start and info.period_end:
        outside = [r for r in good if not info.period_start <= r.txn_date <= info.period_end]
        if outside:
            by_value = [r for r in outside if r.value_date and info.period_start <= r.value_date <= info.period_end]
            other = len(outside) - len(by_value)
            msg = (f"{len(outside)} row(s) are dated just outside the statement period "
                   f"({info.period_start:%d %b %Y} to {info.period_end:%d %b %Y})")
            msg += (" but their value date is inside it (banks list transactions by value date). They are real "
                    "statement entries and are imported; the same entry in a neighbouring statement is recognised "
                    "as a duplicate." if not other else
                    f"; {other} of them also have a value date outside the period. They are imported; check them.")
            checks.append(_check("boundary_rows", "info" if not other else "warn", msg,
                                 rows=[r.row_number for r in outside]))
    return checks


def is_outside_period(r: ParsedRow, info: StatementInfo) -> bool:
    return bool(info.period_start and info.period_end and r.txn_date
                and not info.period_start <= r.txn_date <= info.period_end)


# ------------------------------------------------------------------ privacy for previews

_LONG_DIGITS = re.compile(r"\d{4,}")
_EMAIL = re.compile(r"[\w.+-]+@[\w-]+(\.[\w-]+)+")
HIDDEN = "[hidden]"


def redact_preview_rows(rows: list[list[Any]], header_row: int | None, table_end: int | None, fmt) -> list[list[str]]:
    """Raw rows for the column-mapping screen without the account-holder details.

    Rows above a detected header (name, address, account and customer numbers...) are hidden.
    Footer rows after the table, or every row when no header was detected, have long digit
    runs and e-mail addresses masked, so the layout stays recognisable for choosing the header.
    """
    out = []
    for i, row in enumerate(rows):
        cells = [fmt(c) for c in row]
        if header_row is not None and i < header_row:
            cells = [HIDDEN if c.strip() else "" for c in cells]
        elif header_row is None or (table_end is not None and i >= table_end):
            cells = [_EMAIL.sub("[email]", _LONG_DIGITS.sub(lambda m: "#" * len(m.group()), c)) for c in cells]
        out.append(cells)
    return out
