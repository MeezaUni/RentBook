"""Monthly reports and CSV / XLSX exports (openpyxl, MIT)."""

from __future__ import annotations

import csv
import datetime as dt
import io
import tempfile
import zipfile
from pathlib import Path

from fastapi import APIRouter, Depends
from fastapi.responses import Response
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..db import get_db
from ..errors import DomainError
from ..models import RentDue, RentPayment, Tenancy
from ..money import to_rupees_decimal
from ..periods import current_period, first_day, last_day, validate_period
from ..security import require_user
from ..services import backup as backup_svc
from ..services import cash as cash_svc
from ..services import maintenance as mnt
from ..services import rent as rent_svc
from .dashboard import collected_for_period_by_method, received_in_month
from .properties import buildings_overview
from .rent import ledger

router = APIRouter(prefix="/api/reports", dependencies=[Depends(require_user)])

Table = tuple[list[str], list[list]]


def R(paise: int | None):
    return None if paise is None else to_rupees_decimal(paise)


def monthly_tables(db: Session, period: str) -> dict[str, Table]:
    led = ledger(db, period)
    start, end = first_day(period), last_day(period)
    t: dict[str, Table] = {}
    s = led["summary"]
    recv = received_in_month(db, period)
    by_method = collected_for_period_by_method(db, period)
    m = mnt.month_summary(db, period)
    sal = mnt.salary_month(db, period)
    cb = cash_svc.cashbook(db, period)
    t["summary"] = (["Item", "Amount (₹)", "Notes"], [
        ["Reporting period", period, f"{start} to {end}"],
        ["Expected rent", R(s["expected_paise"]), "Rent due for this month"],
        ["Confirmed rent received for this month", R(s["confirmed_paise"]), "Confirmed payments only"],
        ["  of which cash", R(by_method["cash_paise"]), ""],
        ["  of which online", R(by_method["online_paise"]), ""],
        ["Pending verification", R(s["pending_paise"]), "Not counted as received"],
        ["Outstanding for this month", R(s["outstanding_paise"]), ""],
        ["Overdue for this month", R(s["overdue_paise"]), ""],
        ["Arrears from earlier months", R(s["arrears_paise"]), ""],
        ["Cash rent received during month (all months)", R(recv["cash_confirmed_paise"]), "By receipt date"],
        ["Online rent received during month (all months)", R(recv["online_confirmed_paise"]), "By receipt date"],
        ["Maintenance target", R(m["target_paise"]), ""],
        ["Caretaker cash at start of month", R(m["opening_paise"]), ""],
        ["Maintenance top-up required", R(m["top_up_required_paise"]), "max(0, target - opening)"],
        ["Maintenance cash handed over", R(m["handed_over_paise"]), ""],
        ["Maintenance expenses (total)", R(m["total_expenses_paise"]), ""],
        ["  paid from caretaker cash", R(m["caretaker_spent_paise"]), ""],
        ["  paid directly by family", R(m["family_paid_paise"]), ""],
        ["Caretaker cash at end of month", R(m["closing_paise"]), ""],
        ["Caretaker salary expected", R(sal["expected_paise"]), ""],
        ["Caretaker salary paid (confirmed)", R(sal["confirmed_paise"]), sal["status"]],
        ["Family cash at end of month", R(cb["closing_paise"]), ""],
    ])
    t["rent_dues"] = (["Building", "Flat", "Tenant", "Month", "Type", "Due date", "Rent due (₹)", "Confirmed (₹)",
                       "Pending (₹)", "Outstanding (₹)", "Status"],
                      [[d["building_name"], d["flat_label"], d["tenant_name"], d["period"], d["kind"], d["due_date"],
                        R(d["amount_paise"]), R(d["confirmed_paise"]), R(d["pending_paise"]),
                        R(d["outstanding_paise"]), d["status"]] for d in led["dues"]])
    t["arrears"] = (["Building", "Flat", "Tenant", "Month", "Due date", "Outstanding (₹)"],
                    [[d["building_name"], d["flat_label"], d["tenant_name"], d["period"], d["due_date"],
                      R(d["outstanding_paise"])] for d in led["arrears"]])
    pays = [rent_svc.describe_payment(p) for p in db.scalars(
        select(RentPayment).where(RentPayment.received_date.between(start, end))
        .order_by(RentPayment.received_date, RentPayment.id))]
    t["receipts"] = (["Payment #", "Received", "Building", "Flat", "Tenant", "Method", "Amount (₹)", "For months",
                      "Late", "Status", "Reference", "Recorded by", "Status reason"],
                     [[p["id"], p["received_date"], p["building_name"], p["flat_label"], p["tenant_name"], p["method"],
                       R(p["amount_paise"]), ", ".join(a["period"] for a in p["allocations"]),
                       "yes" if any(a["is_late"] for a in p["allocations"]) else "", p["status"], p["reference"],
                       p["recorded_by"], p["status_reason"]] for p in pays])
    late = [(p, a) for p in pays for a in p["allocations"] if a["is_late"]]
    t["late_receipts"] = (["Payment #", "Received", "Flat", "Tenant", "For month", "Due date", "Amount (₹)", "Status"],
                          [[p["id"], p["received_date"], p["flat_label"], p["tenant_name"], a["period"], a["due_date"],
                            R(a["amount_paise"]), p["status"]] for p, a in late])
    t["buildings"] = (["Building", "Flats", "Expected (₹)", "Confirmed (₹)", "Outstanding (₹)"],
                      [[b["name"], len(b["flats"]), R(sum(f["period_due_paise"] for f in b["flats"])),
                        R(sum(f["period_confirmed_paise"] for f in b["flats"])),
                        R(sum(f["period_outstanding_paise"] for f in b["flats"]))]
                       for b in buildings_overview(db, period)])
    t["maintenance_expenses"] = (["Date", "Category", "Description", "Total (₹)", "From caretaker cash (₹)",
                                  "Paid by family (₹)", "Family method", "Status", "Recorded by"],
                                 [[e["expense_date"], e["category"], e["description"], R(e["total_paise"]),
                                   R(e["caretaker_paid_paise"]), R(e["family_paid_paise"]), e["family_payment_method"],
                                   e["status"], e["recorded_by"]] for e in m["expenses"]])
    t["maintenance_cash"] = (["Date", "Type", "From", "To", "Amount (₹)", "Description", "Status"],
                             [[x["movement_date"], x["kind_label"], x["from_account"], x["to_account"],
                               R(x["amount_paise"]), x["description"], x["status"]] for x in m["movements"]])
    t["salary"] = (["Salary month", "Payment date", "Amount (₹)", "Method", "Reference", "Status", "Notes"],
                   [[p["salary_month"], p["payment_date"], R(p["amount_paise"]), p["method"], p["bank_reference"],
                     p["status"], p["notes"]] for p in sal["payments"]])
    t["cash_reconciliation"] = (["Item", "Amount (₹)"], [
        ["Family cash at start", R(cb["opening_paise"])],
        ["Cash rent received (confirmed)", R(cb["cash_rent_received_paise"])],
        ["Other cash in", R(cb["other_cash_in_paise"])],
        ["Handed to caretaker", R(cb["handed_to_caretaker_paise"])],
        ["Other cash out", R(cb["other_cash_out_paise"])],
        ["Maintenance paid by family in cash", R(cb["family_paid_maintenance_cash_paise"])],
        ["Family cash at end", R(cb["closing_paise"])],
        *[[f"DISCREPANCY: {d['message']}", R(d["amount_paise"])] for d in cb["discrepancies"]],
    ])
    return t


def tenant_history_table(db: Session, tenancy_id: int) -> Table:
    t = db.get(Tenancy, tenancy_id)
    if t is None:
        raise DomainError("not_found", "Tenancy not found", status=404)
    dues = rent_svc.describe_dues(db, list(db.scalars(select(RentDue).where(RentDue.tenancy_id == t.id)
                                                      .order_by(RentDue.period))))
    rows = []
    for d in dues:
        rows.append([d["period"], d["kind"], d["due_date"], R(d["amount_paise"]), R(d["confirmed_paise"]),
                     R(d["pending_paise"]), R(d["outstanding_paise"]), d["status"]])
    return (["Month", "Type", "Due date", "Due (₹)", "Confirmed (₹)", "Pending (₹)", "Outstanding (₹)", "Status"], rows)


def to_csv(table: Table, title: str) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([title])
    w.writerow([f"Generated {dt.datetime.now():%Y-%m-%d %H:%M}. Amounts in rupees. Only 'confirmed' counts as received."])
    w.writerow([])
    w.writerow(table[0])
    w.writerows(table[1])
    return buf.getvalue().encode("utf-8-sig")


def to_xlsx(tables: dict[str, Table], title: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    wb.remove(wb.active)
    for name, (headers, rows) in tables.items():
        ws = wb.create_sheet(name[:31])
        ws.append([f"{title} — {name.replace('_', ' ')}"])
        ws["A1"].font = Font(bold=True, size=12)
        ws.append([f"Generated {dt.datetime.now():%Y-%m-%d %H:%M}. Only 'confirmed' transactions count as received."])
        ws.append([])
        ws.append(headers)
        for c in ws[4]:
            c.font = Font(bold=True)
        for r in rows:
            ws.append(r)
        for i, h in enumerate(headers, start=1):
            if "₹" in h:
                for row in ws.iter_rows(min_row=5, min_col=i, max_col=i):
                    for c in row:
                        c.number_format = "#,##,##0.00"
            width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows[:500] if i - 1 < len(r) and r[i - 1] is not None])
            ws.column_dimensions[ws.cell(row=4, column=i).column_letter].width = min(max(10, width + 2), 60)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _download(data: bytes, filename: str, media: str) -> Response:
    return Response(data, media_type=media, headers={"Content-Disposition": f'attachment; filename="{filename}"'})


XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@router.get("/monthly")
def monthly(period: str | None = None, db: Session = Depends(get_db)):
    period = validate_period(period or current_period())
    tables = monthly_tables(db, period)
    db.commit()
    return {"period": period, "tables": {k: {"headers": h, "rows": [[str(c) if c is not None else "" for c in r]
                                                                      for r in rows]}
                                         for k, (h, rows) in tables.items()}}


@router.get("/export")
def export(period: str | None = None, report: str = "all", format: str = "xlsx", tenancy_id: int | None = None,
           db: Session = Depends(get_db)):
    period = validate_period(period or current_period())
    if report == "tenant_history":
        if not tenancy_id:
            raise DomainError("tenancy_required", "Choose a tenancy")
        tables = {"tenant_history": tenant_history_table(db, tenancy_id)}
        title = f"Tenant history (tenancy {tenancy_id})"
        stem = f"tenant-history-{tenancy_id}"
    else:
        tables = monthly_tables(db, period)
        title = f"Monthly report {period}"
        stem = f"rentbook-{period}"
        if report != "all":
            if report not in tables:
                raise DomainError("unknown_report", f"Unknown report '{report}'")
            tables = {report: tables[report]}
            stem += f"-{report}"
    db.commit()
    if format == "csv":
        if len(tables) != 1:
            raise DomainError("one_table", "CSV exports one report at a time. Choose a report or use XLSX.")
        (name, table), = tables.items()
        return _download(to_csv(table, title), f"{stem}.csv", "text/csv; charset=utf-8")
    if format == "xlsx":
        return _download(to_xlsx(tables, title), f"{stem}.xlsx", XLSX)
    raise DomainError("unknown_format", "Format must be csv or xlsx")


@router.get("/full-export")
def full_export():
    """All essential records as plain CSV files in a zip, readable without this app."""
    with tempfile.TemporaryDirectory() as tmp:
        snap = Path(tmp) / "snap.db"
        backup_svc._sqlite_snapshot(snap)
        files = backup_svc._csv_exports(snap)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("README.txt", "Rentbook full data export. Money columns ending in _paise are in paise "
                                 "(divide by 100 for rupees). Status columns show confirmed / pending / reversed.\n")
        for name, data in files.items():
            z.writestr(name, data)
    return _download(buf.getvalue(), f"rentbook-records-{dt.date.today()}.zip", "application/zip")
